#!/usr/bin/env python3
"""
Build the all-occupations dataset (phase 1: data only, nothing imported by
the app yet) from three official files committed in the repo:

  OEWS May 2025   data/national_M2025_dl.xlsx          (employment, pay, hierarchy)
  EP 2025-35      data/bls_ep_occupation_2025_35.xlsx  (Table 1.2 projections)
  O*NET 30.1      data/db_30_1_text/                   (description, tasks)

Writes src/data/occupations/index.json (small, for the map) and one
src/data/occupations/<major>.json per BLS SOC major group ("continent").

Dev-only tool (openpyxl); nothing here ships in the app bundle. OEWS cells are
parsed with the same code as scripts/extract_oews_wages.py (imported, not
copied), so the symbol rules cannot drift apart.

HARD RULES
  - Every source's SHA-256 must match the constant below, or nothing is read.
  - Values are copied exactly. The only transformation is EP employment, which
    BLS publishes in thousands and src/data/bls_ep_2025_2035.json already
    stores in persons: it is multiplied by 1000 in exact decimal arithmetic
    (aborting if the result is not a whole number). Nothing is rounded,
    estimated, interpolated, split, or borrowed from another code or year.
  - OEWS, EP and O*NET are joined on the exact 7-character SOC only.
  - Every null carries a reason in nullReasons; every value carries the cell
    or line it came from in sources.
  - Output is deterministic: re-running produces byte-identical files.
    EXTRACTED is a constant for that reason; bump it when a source changes.
  - No AI scores. Every occupation is "not yet scored".

O*NET task rule
  O*NET 30.1 has a .00 row in Occupation Data for every SOC it covers, but for
  some "All Other" SOCs the .00 row has no task statements and the tasks sit
  on the sub-codes (.01, .02, ...). So:
    .00 has task statements      -> those tasks only
    otherwise, sub-codes do      -> their tasks, each with its own O*NET code,
                                    identical text kept once (alsoIn lists the
                                    other codes that carry it)
    otherwise                    -> tasks = [] with the reason
  The description always comes from the .00 row when it exists.

USAGE
  python3 scripts/build_occupations.py [--check]
  --check rebuilds in memory and exits 1 if any committed file would change.
"""

import csv
import importlib.util
import json
import os
import re
import sys
from collections import defaultdict
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Reuse the OEWS parser rather than writing a second one.
_spec = importlib.util.spec_from_file_location("extract_oews_wages", os.path.join(ROOT, "scripts", "extract_oews_wages.py"))
oews = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(oews)

import openpyxl  # noqa: E402  (extract_oews_wages already checked it is installed)
from openpyxl.utils import get_column_letter  # noqa: E402

EXTRACTED = "2026-10-24"
OUT_DIR = os.path.join(ROOT, "src", "data", "occupations")
INDEX_BUDGET_BYTES = 250_000
AI_STATUS = "not yet scored"

EP_FILE = os.path.join(ROOT, "data", "bls_ep_occupation_2025_35.xlsx")
EP_SHA256 = "a7d060f0576f7d4f5b829b9a0bb4a6c11a12766d159ebb0e07aab5bd2e7583c5"
EP_SHEET = "Table 1.2"
EP_WAYBACK_CAPTURE = "20261002185424"
EP_ORIGINAL_URL = "https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx"
EP_WAYBACK_URL = f"https://web.archive.org/web/{EP_WAYBACK_CAPTURE}id_/{EP_ORIGINAL_URL}"
EP_SITE_JSON = os.path.join(ROOT, "src", "data", "bls_ep_2025_2035.json")
# JSON field -> Table 1.2 header text. Employment fields are in thousands.
EP_COLUMNS = {
    "employment2025": "Employment, 2025",
    "employment2035": "Employment, 2035",
    "numericChange": "Employment change, numeric, 2025–35",
    "percentChange": "Employment change, percent, 2025–35",
}
EP_THOUSANDS = {"employment2025", "employment2035", "numericChange"}

ONET_DIR = os.path.join(ROOT, "data", "db_30_1_text")
ONET_URL = "https://www.onetcenter.org/dl_files/database/db_30_1_text.zip"
ONET_FILES = {
    "Occupation Data.txt": "63e6029d3d30ff5c7cf39b5304a733b77a409e01c65ae7095fe92f6d18d74a66",
    "Task Statements.txt": "093fae042127cddb2913f78e54c62a7ef6b7a82578d6ea1b60f928f6893470a7",
    "Read Me.txt": None,  # version banner only; checked by content below
}

KELLEY_CSV = os.path.join(ROOT, "data", "Kelley_Job_Map.csv")

# JSON wage field -> OEWS column (same names extract_oews_wages.py uses).
WAGES = oews.MEASURES


def die(msg):
    sys.exit(f"\nbuild_occupations: ABORT: {msg}\n")


def check_hash(path, expected):
    got = oews.sha256(path)
    if got != expected:
        die(f"SHA-256 mismatch for {os.path.relpath(path, ROOT)}\n  expected {expected}\n  got      {got}")
    return got


def thousands_to_persons(v, where):
    persons = Decimal(str(v)) * 1000
    if persons != persons.to_integral_value():
        die(f"{where}: {v} thousand is not a whole number of persons")
    return int(persons)


# ── OEWS ──────────────────────────────────────────────────────────────────

def read_oews():
    wb = openpyxl.load_workbook(oews.WORKBOOK, read_only=True, data_only=True)
    notes_title, fields, notes, topcode = oews.read_notes(wb)
    emp_note = None
    for (a, *_rest) in wb[oews.NOTES_SHEET].iter_rows(values_only=True):
        if isinstance(a, str) and re.match(r"^\*\*\s+=", a.strip()):
            emp_note = a.strip()
    if not emp_note:
        die('OEWS notes sheet has no "**" definition')

    ws = wb[oews.DATA_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    col = {name: i for i, name in enumerate(header)}
    letter = {name: get_column_letter(i + 1) for name, i in col.items()}
    ctx = {"topcode": topcode, "notes": notes, "fields": fields}
    oews.self_test(ctx, ws, col)

    by_group = defaultdict(list)
    for excel_row, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if str(r[col["AREA"]]) != "99" or str(r[col["NAICS"]]) != "000000" or r[col["I_GROUP"]] != "cross-industry":
            continue
        by_group[r[col["O_GROUP"]]].append((excel_row, r))
    return {
        "col": col, "letter": letter, "ctx": ctx, "by_group": by_group,
        "notes_title": notes_title, "topcode": topcode, "fields": fields,
        "not_available_note": notes["*"], "employment_not_available_note": emp_note,
    }


def oews_entry(o, hits):
    """Employment + wages for one detailed code. hits = [(row, values)], detailed row first."""
    col, letter, ctx = o["col"], o["letter"], o["ctx"]
    row, r = hits[0]
    cell = lambda c: f"oews:{oews.DATA_SHEET}!{letter[c]}{row}"  # noqa: E731
    annual_only = r[col["ANNUAL"]] == "TRUE"
    hourly_only = r[col["HOURLY"]] == "TRUE"
    for flag in ("ANNUAL", "HOURLY"):
        if r[col[flag]] not in ("", None, "TRUE"):
            die(f"{r[col['OCC_CODE']]}: unexpected {flag} flag {r[col[flag]]!r}")

    out = {"employment": None}
    sources, null_reasons, topcoded = {}, {}, {}
    emp = r[col["TOT_EMP"]]
    sources["employment"] = cell("TOT_EMP")
    if isinstance(emp, int) and not isinstance(emp, bool):
        out["employment"] = emp
    elif emp == "**":
        null_reasons["employment"] = f'Not available: "**" in the BLS workbook ({o["employment_not_available_note"]})'
    else:
        die(f"{r[col['OCC_CODE']]} TOT_EMP (row {row}): unrecognised value {emp!r}")

    for m, c in WAGES.items():
        try:
            value, top, reason = oews.read_cell(r[col[c]], m, annual_only, hourly_only, ctx)
        except ValueError as err:
            die(f"{r[col['OCC_CODE']]} {c} (row {row}): {err}")
        out[m] = value
        sources[m] = cell(c)
        if top:
            topcoded[m] = True
        if reason:
            null_reasons[m] = reason
    out["annualOnly"] = annual_only
    out["hourlyOnly"] = hourly_only
    sources["annualOnly"] = cell("ANNUAL")
    sources["hourlyOnly"] = cell("HOURLY")
    return out, sources, null_reasons, topcoded


# ── EP ────────────────────────────────────────────────────────────────────

def read_ep():
    wb = openpyxl.load_workbook(EP_FILE, read_only=True, data_only=True)
    ws = wb[EP_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    title = rows[0][0]
    if "2025–35" not in str(title):
        die(f'EP sheet title is "{title}", expected 2025–35 projections')
    header = [str(h).strip() if h is not None else None for h in rows[1]]
    col = {}
    for name in ["2025 National Employment Matrix title", "2025 National Employment Matrix code", "Occupation type", *EP_COLUMNS.values()]:
        if header.count(name) != 1:
            die(f'EP header "{name}" found {header.count(name)} times')
        col[name] = header.index(name)
    code_c, type_c, title_c = col["2025 National Employment Matrix code"], col["Occupation type"], col["2025 National Employment Matrix title"]

    lines, summaries = {}, set()
    for excel_row, r in enumerate(rows[2:], start=3):
        kind = r[type_c]
        if kind == "Summary":
            summaries.add(r[code_c])
        if kind != "Line item":
            continue
        soc = r[code_c]
        if soc in lines:
            die(f"EP line item {soc} appears twice")
        entry = {"title": r[title_c].strip()}
        sources = {"title": f"ep:{EP_SHEET}!{get_column_letter(title_c + 1)}{excel_row}"}
        for field, name in EP_COLUMNS.items():
            v = r[col[name]]
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                die(f"EP {soc} {name} (row {excel_row}): non-numeric {v!r}")
            entry[field] = thousands_to_persons(v, f"EP {soc} {name}") if field in EP_THOUSANDS else v
            sources[field] = f"ep:{EP_SHEET}!{get_column_letter(col[name] + 1)}{excel_row}"
        lines[soc] = (entry, sources)
    return {"title": title, "lines": lines, "summaries": summaries}


# ── O*NET ─────────────────────────────────────────────────────────────────

def read_tsv(name):
    """Tab-separated O*NET file -> [(line_number, {column: value})]. No quote handling: text is kept byte-exact."""
    with open(os.path.join(ONET_DIR, name), encoding="utf-8", newline="") as f:
        reader = csv.reader(f, delimiter="\t", quoting=csv.QUOTE_NONE)
        header = next(reader)
        out = []
        for line, r in enumerate(reader, start=2):
            if len(r) != len(header):
                die(f"{name} line {line}: {len(r)} fields, expected {len(header)}")
            out.append((line, dict(zip(header, r))))
    return out


def read_onet():
    readme = open(os.path.join(ONET_DIR, "Read Me.txt"), encoding="utf-8").read()
    if not readme.startswith("O*NET 30.1 Database"):
        die("data/db_30_1_text/Read Me.txt is not O*NET 30.1")
    release = readme.splitlines()[1].strip()
    hashes = {n: check_hash(os.path.join(ONET_DIR, n), h) for n, h in ONET_FILES.items() if h}

    occ = {}
    for line, r in read_tsv("Occupation Data.txt"):
        occ[r["O*NET-SOC Code"]] = (line, r)
    tasks = defaultdict(list)
    for line, r in read_tsv("Task Statements.txt"):
        if r["O*NET-SOC Code"] not in occ:
            die(f"Task Statements.txt line {line}: {r['O*NET-SOC Code']} not in Occupation Data.txt")
        tasks[r["O*NET-SOC Code"]].append((line, r))
    by_soc = defaultdict(list)
    for code in occ:
        by_soc[code[:7]].append(code)
    return {"occ": occ, "tasks": tasks, "by_soc": {k: sorted(v) for k, v in by_soc.items()}, "release": release, "hashes": hashes}


def task_json(line, r, also_in=None):
    t = {"id": int(r["Task ID"]), "text": r["Task"], "type": r["Task Type"], "onetCode": r["O*NET-SOC Code"], "line": line}
    if also_in:
        t["alsoIn"] = also_in
    return t


def onet_entry(on, soc):
    codes = on["by_soc"].get(soc, [])
    base = soc + ".00"
    out = {"onetCode": None, "onetTaskSource": "none", "onetSubCodes": [], "description": None, "tasks": [], "dedupedTasks": 0}
    sources, null_reasons = {}, {}

    if base in on["occ"]:
        line, r = on["occ"][base]
        out["onetCode"] = base
        out["description"] = r["Description"]
        sources["onetCode"] = f"onet:Occupation Data.txt#L{line}:O*NET-SOC Code"
        sources["description"] = f"onet:Occupation Data.txt#L{line}:Description"
    else:
        null_reasons["description"] = f"No O*NET 30.1 occupation code starts with {soc}"
        null_reasons["onetCode"] = null_reasons["description"]

    subs = [c for c in codes if c != base]
    if on["tasks"].get(base):
        out["onetTaskSource"] = ".00"
        out["tasks"] = [task_json(line, r) for line, r in on["tasks"][base]]
    elif any(on["tasks"].get(c) for c in subs):
        out["onetTaskSource"] = "sub-codes"
        seen, kept = {}, []
        for c in subs:
            for line, r in on["tasks"].get(c, []):
                if r["Task"] in seen:
                    seen[r["Task"]].setdefault("alsoIn", []).append(c)
                    out["dedupedTasks"] += 1
                    continue
                t = task_json(line, r)
                seen[r["Task"]] = t
                kept.append(t)
        out["tasks"] = kept
    else:
        if not codes:
            why = f"No O*NET 30.1 occupation code starts with {soc}"
            if soc.endswith("0"):
                group = sorted(c for c in on["occ"] if c[:6] == soc[:6])
                if group:
                    why += f"; O*NET lists this BLS group's members separately ({', '.join(group)}) and they are not combined here"
        elif subs:
            why = f"O*NET 30.1 has no task statements for {', '.join(codes)}"
        else:
            why = f"O*NET 30.1 has no task statements for {base}"
        null_reasons["tasks"] = why
    out["onetSubCodes"] = [{"code": c, "title": on["occ"][c][1]["Title"], "taskCount": len(on["tasks"].get(c, []))} for c in subs]
    return out, sources, null_reasons


# ── Kelley lens ───────────────────────────────────────────────────────────

def kelley_lens():
    """SOC -> {tracks, titles}. Titles join to the site's SOC (MAP_TITLE_TO_SOC); the CSV supplies the track."""
    site = oews.soc_map()
    with open(KELLEY_CSV, encoding="utf-8", newline="") as f:
        lines = f.read().splitlines()
    header = next(csv.reader([lines[0].strip('"')]))
    rows = [dict(zip(header, next(csv.reader([ln.strip('"')])))) for ln in lines[1:] if ln.strip()]
    track = {r["kelley_title"]: r["track"] for r in rows}
    if set(track) != set(site):
        die(f"Kelley_Job_Map.csv titles differ from MAP_TITLE_TO_SOC: {sorted(set(track) ^ set(site))}")
    csv_soc_differs = [
        {"title": r["kelley_title"], "csvSoc": r["soc_code"], "siteSoc": site[r["kelley_title"]]}
        for r in rows if r["soc_code"] != site[r["kelley_title"]]
    ]
    lens = defaultdict(lambda: {"tracks": set(), "titles": []})
    for title in sorted(site):
        lens[site[title]]["tracks"].add(track[title])
        lens[site[title]]["titles"].append(title)
    return {s: {"tracks": sorted(v["tracks"]), "titles": v["titles"]} for s, v in lens.items()}, csv_soc_differs


# ── Build ─────────────────────────────────────────────────────────────────

def build():
    print("\nAll-occupations dataset\n")
    oews_hash = check_hash(oews.WORKBOOK, oews.EXPECTED_SHA256)
    ep_hash = check_hash(EP_FILE, EP_SHA256)
    print(f"  OEWS  sha256 OK {oews_hash}")
    print(f"  EP    sha256 OK {ep_hash}")
    o = read_oews()
    ep = read_ep()
    on = read_onet()
    for n, h in on["hashes"].items():
        print(f"  O*NET sha256 OK {h}  {n}")
    lens, csv_soc_differs = kelley_lens()

    col = o["col"]
    code = lambda r: r[col["OCC_CODE"]]  # noqa: E731
    title = lambda r: r[col["OCC_TITLE"]]  # noqa: E731
    emp_cell = lambda row: f"oews:{oews.DATA_SHEET}!{o['letter']['TOT_EMP']}{row}"  # noqa: E731
    title_cell = lambda row: f"oews:{oews.DATA_SHEET}!{o['letter']['OCC_TITLE']}{row}"  # noqa: E731

    majors = sorted(o["by_group"]["major"], key=lambda h: code(h[1]))
    minors = sorted(o["by_group"]["minor"], key=lambda h: code(h[1]))
    broads = defaultdict(list)
    for row, r in o["by_group"]["broad"]:
        broads[code(r)].append((row, r))
    detailed = defaultdict(list)
    for row, r in o["by_group"]["detailed"]:
        detailed[code(r)].append((row, r))
    if len(majors) != 22:
        die(f"expected 22 SOC major groups, found {len(majors)}")
    for soc, hits in detailed.items():
        if len(hits) != 1:
            die(f"detailed code {soc} appears {len(hits)} times")
        if soc in broads:  # OEWS publishes this broad code as its most detailed level
            hits.append(broads[soc][0])

    def minor_of(soc):
        cands = [code(r) for _, r in minors if code(r)[:2] == soc[:2] and code(r) <= soc]
        if not cands:
            die(f"{soc} has no minor group")
        return cands[-1]

    def broad_of(soc):
        """SOC broad code (first five digits), or None when OEWS publishes no row for it."""
        b = soc[:6] + "0"
        return b if b in broads else None

    for b in broads:
        if b[:2] not in {code(r)[:2] for _, r in majors}:
            die(f"broad {b} has no major group")

    # Cross-dataset code lists (exact SOC only, never forced).
    oews_codes = set(detailed)
    unmatched = {
        "oewsWithoutEp": [{"soc": s, "title": title(detailed[s][0][1])} for s in sorted(oews_codes - set(ep["lines"]))],
        "epWithoutOews": [{"soc": s, "title": ep["lines"][s][0]["title"]} for s in sorted(set(ep["lines"]) - oews_codes)],
        "oewsWithoutOnetCode": [{"soc": s, "title": title(detailed[s][0][1])} for s in sorted(oews_codes) if s not in on["by_soc"]],
        "onetSocsWithoutOews": [
            {"soc": s, "onetCodes": on["by_soc"][s]} for s in sorted(on["by_soc"]) if s not in oews_codes
        ],
    }

    continents, index_occs, files = [], [], {}
    for mrow, mr in majors:
        major = code(mr)
        mm = major[:2]
        entries, minor_groups, broad_groups = {}, [], []
        for nrow, nr in minors:
            if code(nr)[:2] != mm:
                continue
            minor_groups.append({"soc": code(nr), "title": title(nr), "employment": nr[col["TOT_EMP"]],
                                 "sources": {"employment": emp_cell(nrow), "title": title_cell(nrow)}, "socs": []})
        for b in sorted(broads):
            if b[:2] == mm:
                brow, br = broads[b][0]
                broad_groups.append({"soc": b, "title": title(br), "minor": minor_of(b), "employment": br[col["TOT_EMP"]],
                                     "sources": {"employment": emp_cell(brow), "title": title_cell(brow)}})
        for soc in sorted(s for s in detailed if s[:2] == mm):
            hits = detailed[soc]
            r = hits[0][1]
            minor = minor_of(soc)
            values, sources, null_reasons, topcoded = oews_entry(o, hits)
            sources["title"] = title_cell(hits[0][0])

            if soc in ep["lines"]:
                ep_vals, ep_src = ep["lines"][soc]
                ep_block = dict(ep_vals)
                sources.update({f"ep.{k}": v for k, v in ep_src.items()})
            else:
                ep_block = None
                null_reasons["ep"] = f'No "Line item" row for {soc} in BLS EP {EP_SHEET} (2025–35)'

            broad = broad_of(soc)
            if broad is None:
                null_reasons["broad"] = f"OEWS May 2025 publishes no broad-level row for {soc[:6]}0"

            onet_vals, onet_src, onet_null = onet_entry(on, soc)
            sources.update(onet_src)
            null_reasons.update(onet_null)

            entry = {
                "soc": soc,
                "title": title(r),
                "major": major,
                "minor": minor,
                "broad": broad,
                "oewsRows": [h[0] for h in hits],
                "oewsOGroups": [h[1][col["O_GROUP"]] for h in hits],
                **values,
                "topcoded": topcoded,
                "ep": ep_block,
                **onet_vals,
                "aiStatus": AI_STATUS,
            }
            if soc in lens:
                entry["kelleyTracks"] = lens[soc]["tracks"]
                entry["kelleyTitles"] = lens[soc]["titles"]
            entry["sources"] = sources
            entry["nullReasons"] = null_reasons
            entries[soc] = entry
            next(g for g in minor_groups if g["soc"] == minor)["socs"].append(soc)

            row = {
                "soc": soc,
                "title": title(r),
                "minor": minor,
                "employment": values["employment"],
                "annualMedian": values["annualMedian"],
                "epPercentChange": ep_block["percentChange"] if ep_block else None,
                "taskCount": len(onet_vals["tasks"]),
            }
            if soc in lens:
                row["kelleyTracks"] = lens[soc]["tracks"]
            index_occs.append(row)

        continent = {
            "soc": major,
            "code": mm,
            "title": title(mr),
            "employment": mr[col["TOT_EMP"]],
            "sources": {"employment": emp_cell(mrow), "title": title_cell(mrow)},
            "file": f"{mm}.json",
            "occupationCount": len(entries),
            "minorGroups": minor_groups,
        }
        continents.append(continent)
        files[f"{mm}.json"] = {
            "_meta": {
                "continent": major,
                "title": title(mr),
                "sources": "see index.json _meta.sources; source refs are <source>:<sheet>!<cell> or <source>:<file>#L<line>:<column>",
                "extracted": EXTRACTED,
                "generator": "python3 scripts/build_occupations.py",
            },
            "broadGroups": broad_groups,
            "occupations": entries,
        }

    total_row, total_r = o["by_group"]["total"][0]
    meta = {
        "description": "Every BLS detailed occupation, grouped by the 22 SOC major groups. Phase 1: data only, not imported by the app.",
        "extracted": EXTRACTED,
        "generator": "python3 scripts/build_occupations.py",
        "aiStatus": f'All occupations "{AI_STATUS}". The 50 Kelley roles keep their published scores in src/data/ai_scores.json, which this dataset does not copy.',
        "sources": {
            "oews": {
                "name": "BLS Occupational Employment and Wage Statistics, national",
                "vintage": oews.BLS_RELEASE,
                "file": "data/national_M2025_dl.xlsx",
                "url": oews.WORKBOOK_URL,
                "sha256": oews_hash,
                "sheet": oews.DATA_SHEET,
                "filter": "AREA = 99 (U.S.), NAICS = 000000, I_GROUP = cross-industry",
                "workbookNotesTitle": o["notes_title"],
                "topcode": o["topcode"],
                "notAvailableNote": o["not_available_note"],
                "employmentNotAvailableNote": o["employment_not_available_note"],
                "columns": {"employment": "TOT_EMP", **WAGES, "annualOnly": "ANNUAL", "hourlyOnly": "HOURLY"},
                "fieldDefinitions": {c: o["fields"].get(c) for c in ["TOT_EMP", *WAGES.values(), "ANNUAL", "HOURLY", "O_GROUP"]},
                "allOccupations": {"employment": total_r[col["TOT_EMP"]], "sources": {"employment": emp_cell(total_row)}},
            },
            "ep": {
                "name": "BLS Employment Projections, Table 1.2 Occupational projections",
                "vintage": "2025–2035",
                "file": "data/bls_ep_occupation_2025_35.xlsx",
                "url": EP_ORIGINAL_URL,
                "waybackUrl": EP_WAYBACK_URL,
                "waybackCapture": EP_WAYBACK_CAPTURE,
                "sha256": ep_hash,
                "sheet": EP_SHEET,
                "sheetTitle": ep["title"],
                "rows": 'Occupation type = "Line item" only',
                "columns": EP_COLUMNS,
                "units": "employment2025, employment2035 and numericChange are persons: BLS thousands × 1000 in exact decimal "
                         "arithmetic, the same units as src/data/bls_ep_2025_2035.json. percentChange is copied as published.",
            },
            "onet": {
                "name": "O*NET Database",
                "vintage": "O*NET 30.1",
                "release": on["release"],
                "dir": "data/db_30_1_text/",
                "url": ONET_URL,
                "sha256": on["hashes"],
                "taskRule": "Tasks from the .00 code when it has task statements; otherwise from its sub-codes (each task keeps its "
                            "own O*NET code; identical text kept once, alsoIn lists the other codes). Description from the .00 row.",
            },
            "kelley": {
                "file": "data/Kelley_Job_Map.csv",
                "join": "kelley_title -> MAP_TITLE_TO_SOC (src/utils/onet.ts, the site's SOC) -> kelleyTracks from the CSV's track column",
                "csvSocDiffersFromSite": csv_soc_differs,
            },
        },
        "matching": "OEWS, EP and O*NET joined on the exact SOC code only.",
        "unmatched": unmatched,
        "counts": {
            "continents": len(continents),
            "occupations": len(index_occs),
            "withEmployment": sum(1 for r in index_occs if r["employment"] is not None),
            "withAnnualMedian": sum(1 for r in index_occs if r["annualMedian"] is not None),
            "withEp": sum(1 for r in index_occs if r["epPercentChange"] is not None),
            "withTasks": sum(1 for r in index_occs if r["taskCount"]),
            "tasks": sum(r["taskCount"] for r in index_occs),
        },
    }
    index = {"_meta": meta, "continents": continents, "occupations": index_occs}
    return index, files


def dumps(obj):
    return json.dumps(obj, indent=1, ensure_ascii=False) + "\n"


def main():
    check = "--check" in sys.argv
    index, files = build()
    outputs = {"index.json": dumps(index), **{n: dumps(v) for n, v in files.items()}}
    size = len(outputs["index.json"].encode("utf-8"))
    if size > INDEX_BUDGET_BYTES:
        die(f"index.json is {size:,} bytes, over the {INDEX_BUDGET_BYTES:,} budget")

    changed = []
    for name, text in outputs.items():
        path = os.path.join(OUT_DIR, name)
        old = open(path, encoding="utf-8").read() if os.path.exists(path) else None
        if old != text:
            changed.append(name)
            if not check:
                os.makedirs(OUT_DIR, exist_ok=True)
                with open(path, "w", encoding="utf-8", newline="\n") as f:
                    f.write(text)
    stale = sorted(set(os.listdir(OUT_DIR)) - set(outputs)) if os.path.isdir(OUT_DIR) else []
    if stale:
        die(f"unexpected files in {os.path.relpath(OUT_DIR, ROOT)}: {stale}")

    c = index["_meta"]["counts"]
    print(f"\n  {c['continents']} continents, {c['occupations']} detailed occupations")
    print(f"  employment {c['withEmployment']}, annual median {c['withAnnualMedian']}, EP {c['withEp']}, "
          f"O*NET tasks {c['withTasks']} ({c['tasks']:,} tasks)")
    print(f"  index.json {size:,} bytes (budget {INDEX_BUDGET_BYTES:,})")
    if check:
        print(f"  --check: {len(changed)} file(s) would change{': ' + ', '.join(changed) if changed else ''}")
        sys.exit(1 if changed else 0)
    print(f"  wrote {len(changed)} changed file(s) of {len(outputs)} in {os.path.relpath(OUT_DIR, ROOT)}\n")


if __name__ == "__main__":
    main()
